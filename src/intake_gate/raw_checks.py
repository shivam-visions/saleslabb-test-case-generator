"""Raw-file integrity checks (IPM-facing, carrier-agnostic).

The IPM uploads RAW insurer files (rates / benefits / confluence docs).
No standardized workbooks exist at this stage, so these checks answer the
IPM's actual question: "is what the insurer sent machine-readable and
complete enough to onboard, and is anything obviously missing or broken?"

Deterministic only — no LLM. Categories come from `_categories.json`
(written by the backend next to the materialized files): name -> category
(rates|benefits|confluence|other).
"""
from __future__ import annotations

import csv as csv_mod
import io
import json
import re
import subprocess
from pathlib import Path

from openpyxl import load_workbook

from .models import Severity

NUMERIC_RE = re.compile(r"\d")
ROW_SCAN_LIMIT = 3000
BLANK_EXAMPLE_CAP = 8


def _f(rule, severity, file, message, sheet=None, row=None):
    return {"rule": rule, "severity": severity, "file": file,
            "sheet": sheet, "row": row, "message": message}


def _pdf_text(path: Path) -> tuple[str | None, str | None]:
    """Extract text via poppler pdftotext (same dependency the extractors
    use). Returns (text, error)."""
    try:
        res = subprocess.run(
            ["pdftotext", "-layout", str(path), "-"],
            capture_output=True, timeout=60)
        if res.returncode != 0:
            return None, res.stderr.decode(errors="replace")[:200]
        return res.stdout.decode(errors="replace"), None
    except FileNotFoundError:
        return None, "pdftotext not installed"
    except subprocess.TimeoutExpired:
        return None, "pdftotext timed out"


def _check_pdf(path: Path, category: str, add) -> None:
    text, err = _pdf_text(path)
    if text is None:
        add(_f("RAW-PDF-001", "ERROR", path.name,
               f"PDF could not be read ({err}) — if it is password-protected "
               "or corrupted, request a clean copy from the insurer"))
        return
    stripped = text.strip()
    if len(stripped) < 100:
        add(_f("RAW-PDF-002", "WARN" if category != "rates" else "ERROR",
               path.name,
               "PDF contains (almost) no extractable text — it appears to be "
               "a scanned image. Rates/benefits cannot be machine-read from "
               "scans; request the original file from the insurer"))
        return
    if category == "rates":
        numbers = re.findall(r"\d[\d,]*\.?\d*", stripped)
        if len(numbers) < 20:
            add(_f("RAW-RATE-001", "WARN", path.name,
                   f"only {len(numbers)} numeric values found in a rates PDF "
                   "— verify this is actually the rate sheet"))
        else:
            add(_f("RAW-OK", "INFO", path.name,
                   f"rates PDF readable: ~{len(numbers)} numeric values, "
                   f"{len(stripped.splitlines())} lines extracted"))
    else:
        add(_f("RAW-OK", "INFO", path.name,
               f"readable: {len(stripped.splitlines())} text lines extracted"))


def _scan_numeric_blanks(rows: list[tuple], file: str, sheet: str, add) -> None:
    """Find blank cells inside mostly-numeric columns of a rates table —
    the classic 'empty copay/premium cell in the rate sheet' defect
    (real example: PROD-2179)."""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    numeric_cells = 0
    blanks = []
    for c in range(ncols):
        col = [(i, r[c] if c < len(r) else None) for i, r in enumerate(rows)]
        nums = [i for i, v in col if isinstance(v, (int, float))]
        if len(nums) < 5:
            continue
        if len(nums) / max(1, sum(1 for _, v in col if v not in (None, ""))) < 0.6:
            continue
        numeric_cells += len(nums)
        first, last = min(nums), max(nums)
        for i, v in col:
            if first < i < last and (v is None or v == ""):
                blanks.append((i + 1, c + 1))
    if numeric_cells == 0:
        add(_f("RAW-RATE-002", "WARN", file,
               "no numeric table detected in a rates spreadsheet — verify "
               "this is the rate sheet", sheet=sheet))
        return
    if blanks:
        ex = ", ".join(f"r{r}c{c}" for r, c in blanks[:BLANK_EXAMPLE_CAP])
        add(_f("RAW-RATE-003", "WARN", file,
               f"{len(blanks)} blank cell(s) inside numeric rate columns "
               f"(e.g. {ex}) — empty premiums/copays in the source caused "
               "real production queries before; confirm with the insurer",
               sheet=sheet))
    else:
        add(_f("RAW-OK", "INFO", file,
               f"numeric rate table detected ({numeric_cells} numeric cells, "
               "no blanks inside rate columns)", sheet=sheet))


def _check_xlsx(path: Path, category: str, add) -> None:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        add(_f("RAW-XLS-001", "ERROR", path.name,
               f"spreadsheet could not be opened ({type(e).__name__}) — "
               "request a clean copy"))
        return
    try:
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= ROW_SCAN_LIMIT:
                    break
                if any(v not in (None, "") for v in row):
                    rows.append(row)
            if not rows:
                continue
            if category == "rates":
                _scan_numeric_blanks(rows, path.name, ws.title, add)
            else:
                add(_f("RAW-OK", "INFO", path.name,
                       f"readable: {len(rows)} non-empty rows", sheet=ws.title))
    finally:
        wb.close()


def _check_csv(path: Path, category: str, add) -> None:
    try:
        text = path.read_text(errors="replace")
        rows = list(csv_mod.reader(io.StringIO(text)))
    except Exception as e:
        add(_f("RAW-CSV-001", "ERROR", path.name,
               f"CSV could not be parsed ({type(e).__name__})"))
        return
    body = [r for r in rows if any(c.strip() for c in r)]
    if not body:
        add(_f("RAW-CSV-002", "ERROR", path.name, "CSV is empty"))
        return
    add(_f("RAW-OK", "INFO", path.name, f"readable: {len(body)} rows"))


def run_raw_checks(folder: str | Path) -> dict:
    folder = Path(folder)
    cats_file = folder / "_categories.json"
    categories: dict = {}
    if cats_file.exists():
        categories = json.loads(cats_file.read_text())

    findings: list[dict] = []
    add = findings.append
    files = [p for p in sorted(folder.iterdir())
             if p.is_file() and not p.name.startswith("_")]

    for p in files:
        cat = categories.get(p.name, "other")
        ext = p.suffix.lower()
        if ext == ".pdf":
            _check_pdf(p, cat, add)
        elif ext in (".xlsx", ".xlsm"):
            _check_xlsx(p, cat, add)
        elif ext == ".csv":
            _check_csv(p, cat, add)
        elif ext == ".xls":
            add(_f("RAW-XLS-002", "WARN", p.name,
                   "legacy .xls format — cannot be machine-checked; "
                   "re-save as .xlsx for full validation"))
        # txt/md/doc/img/etc: stored + diffed, no content checks

    sevs = {f["severity"] for f in findings}
    verdict = ("FAIL" if "ERROR" in sevs
               else "PASS_WITH_WARNINGS" if "WARN" in sevs else "PASS")
    return {
        "folder": str(folder),
        "verdict": verdict,
        "files": [p.name for p in files],
        "counts": {
            "error": sum(1 for f in findings if f["severity"] == "ERROR"),
            "warn": sum(1 for f in findings if f["severity"] == "WARN"),
            "info": sum(1 for f in findings if f["severity"] == "INFO"),
        },
        "findings": findings,
    }
